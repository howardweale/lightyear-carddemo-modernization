"""Regenerate the user manual using python-docx, preserving its typography."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]


def inline(p,text):
    for token in re.split(r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))",text):
        if token.startswith("[") and "](" in token:
            label,url=token[1:].split("](",1);url=url[:-1]
            if not url.startswith(("http:","https:","#")):url="https://github.com/howardweale/lightyear-carddemo-modernization/blob/main/docs/"+url
            h=OxmlElement("w:hyperlink");h.set(qn("r:id"),p.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement("w:r");t=OxmlElement("w:t");t.text=label;r.append(t);h.append(r);p._p.append(h)
        else:
            r=p.add_run(token.strip("`") if token.startswith("`") else token[2:-2] if token.startswith("**") else token)
            if token.startswith("**"):r.bold=True
            if token.startswith("`"):r.font.name="Consolas";r.font.size=Pt(9)


def build():
    path=ROOT/"docs/control-tower-users-manual.docx"
    d=Document(path)
    for item in list(d._element.body):
        if item.tag!=qn("w:sectPr"):d._element.body.remove(item)
    sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.left_margin=sec.right_margin=Inches(.8);sec.top_margin=sec.bottom_margin=Inches(.7)
    for name in ("Normal","Title","Heading 1","Heading 2","Heading 3"):
        s=d.styles[name];s.font.name="Calibri";s.font.color.rgb=RGBColor(0,0,0)
        if s._element.rPr is not None:
            for c in s._element.rPr.findall(qn("w:color")):c.attrib.pop(qn("w:themeColor"),None)
    d.styles["Normal"].font.size=Pt(11);d.styles["Normal"].paragraph_format.space_after=Pt(6)
    d.styles["Normal"].paragraph_format.line_spacing=1.06
    lines=(ROOT/"docs/control-tower-users-manual.md").read_text(encoding="utf-8").splitlines();i=0
    while i<len(lines):
        line=lines[i];i+=1
        if not line.strip():continue
        if line.startswith("#"):
            level=len(line)-len(line.lstrip("#"));p=d.add_paragraph(style="Title" if level==1 else "Heading "+str(min(3,level-1)));inline(p,line[level:].strip());continue
        if line.startswith("|"):
            rows=[line]
            while i<len(lines) and lines[i].startswith("|"):rows.append(lines[i]);i+=1
            rows=[r for r in rows if not re.fullmatch(r"[| :\-]+",r)]
            cells=[[v.strip() for v in r.strip("|").split("|")] for r in rows]
            t=d.add_table(rows=0,cols=len(cells[0]));t.autofit=False
            widths=([1.5,2.5,2.9] if len(cells[0])==3 else [1.8,5.1] if len(cells[0])==2 else [6.9/len(cells[0])]*len(cells[0]))
            for col,w in zip(t.columns,widths):col.width=Inches(w)
            borders=OxmlElement("w:tblBorders")
            for edge in ("top","left","bottom","right","insideH","insideV"):
                el=OxmlElement("w:"+edge);el.set(qn("w:val"),"single");el.set(qn("w:sz"),"4");el.set(qn("w:color"),"D9D9D9");borders.append(el)
            t._tbl.tblPr.append(borders)
            for n,row in enumerate(cells):
                tr=t.add_row();pr=tr._tr.get_or_add_trPr();pr.append(OxmlElement("w:cantSplit"))
                if n==0:pr.append(OxmlElement("w:tblHeader"))
                for c,text,w in zip(tr.cells,row,widths):
                    c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tc=c._tc.get_or_add_tcPr();m=OxmlElement("w:tcMar")
                    for side in ("top","left","bottom","right"):
                        v=OxmlElement("w:"+side);v.set(qn("w:w"),"90");v.set(qn("w:type"),"dxa");m.append(v)
                    tc.append(m);shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"E8EFF4" if n==0 else "FFFFFF");tc.append(shade)
                    p=c.paragraphs[0];p.paragraph_format.space_after=Pt(3);inline(p,text)
                    for r in p.runs:
                        r.font.size=Pt(10)
                        if n==0:r.bold=True
            d.add_paragraph();continue
        if line.startswith("```"):
            code=[]
            while i<len(lines) and not lines[i].startswith("```"):code.append(lines[i]);i+=1
            i+=1;p=d.add_paragraph();r=p.add_run("\n".join(code));r.font.name="Consolas";r.font.size=Pt(9);continue
        bullet=re.match(r"^(- |\d+\. )",line)
        prefix=bullet.group() if bullet else ""
        text=line[len(prefix):].removeprefix("> ")
        while i<len(lines) and lines[i].strip() and not re.match(r"^(#|\||```|- |\d+\. )",lines[i]):text+=" "+lines[i].removeprefix("> ").strip();i+=1
        p=d.add_paragraph(style="List Bullet" if prefix=="- " else "Normal");inline(p,(prefix if prefix and prefix!="- " else "")+text)
    d.core_properties.title="Control Tower user manual"
    d.core_properties.subject="Factory scope and scoped signed-decision review"
    d.save(path)
    print(path)


if __name__=="__main__":build()
